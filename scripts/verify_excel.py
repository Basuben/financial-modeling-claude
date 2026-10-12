"""Recalculate the workbook in LibreOffice and compare it with the Python model.

Needs LibreOffice (soffice) on the PATH and a workbook built with
scripts/build_excel.py. Exits with an error if any number differs by more than 0.01.
"""
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
import model  # noqa: E402

xlsx = ROOT / "outputs" / "savanna_foods_model.xlsx"
with tempfile.TemporaryDirectory() as tmp:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", tmp, str(xlsx)],
        check=True, capture_output=True,
    )
    wb = load_workbook(Path(tmp) / xlsx.name, data_only=True)

r = model.run_all()
errors, checked = [], 0


def cmp(label, got, want):
    global checked
    checked += 1
    if got is None or abs(got - want) > 0.01:
        errors.append(f"{label}: excel={got} python={want}")


proj = r["proj"]
rowmap = {4: "revenue", 7: "ebitda", 9: "ebit", 13: "net_income", 22: "cfo", 23: "cash", 25: "ufcf"}
for row, key in rowmap.items():
    for i, col in enumerate("BCDEF"):
        cmp(f"Model {key} {2026 + i}", wb["Model"][f"{col}{row}"].value, proj[key].iloc[i])
for i, col in enumerate("BCDEF"):
    cmp(f"balance check {col}", wb["Model"][f"{col}26"].value, float(proj["balance_check"].iloc[i]))
d = r["dcf"]
for cell, key in {"B8": "pv_fcf", "B9": "terminal_value", "B11": "enterprise_value",
                  "B13": "equity_value", "B14": "per_share"}.items():
    cmp(f"DCF {key}", wb["DCF"][cell].value, d[key])
sens = r["sens"]
for i in range(5):
    for j in range(5):
        cmp(f"Sensitivity {i},{j}", wb["Sensitivity"].cell(row=4 + i, column=2 + j).value, sens.iloc[i, j])

if errors:
    print("MISMATCH")
    print("\n".join(errors))
    sys.exit(1)
print(f"OK: {checked} cells match the Python model within 0.01")
