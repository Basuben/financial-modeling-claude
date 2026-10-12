"""Build outputs/savanna_foods_model.xlsx with live formulas.

Change a blue input on the Assumptions sheet and the projections, DCF and
sensitivity grid all recalculate in Excel. The values match the Python model,
which scripts/verify_excel.py checks using LibreOffice.
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
from model import Assumptions, load_history  # noqa: E402

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDE6F0")
N0, PCT = "#,##0;(#,##0)", "0.0%"
YEARS = [2026, 2027, 2028, 2029, 2030]
COLS = "BCDEF"


def header(ws, row, labels, start_col=2):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=text)
        c.font, c.fill, c.alignment = BOLD, HEAD, Alignment(horizontal="center")


def build(path: Path) -> None:
    a, hist = Assumptions(), load_history().reset_index()
    wb = Workbook()

    # Assumptions
    ws = wb.active
    ws.title = "Assumptions"
    ws["A1"] = "Savanna Foods Ltd (fictional). KES millions. Blue cells are inputs."
    ws["A1"].font = BOLD
    ws["A3"] = "Driver by year"
    header(ws, 3, YEARS, start_col=3)
    for r, (label, vals) in enumerate(
        [("Revenue growth", a.growth), ("COGS % revenue", a.cogs_pct), ("SG&A % revenue", a.sga_pct)], start=4
    ):
        ws.cell(row=r, column=1, value=label)
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=3 + i, value=v)
            c.font, c.number_format = BLUE, PCT
    scalars = [
        ("Depreciation % opening PPE", a.dep_rate, PCT), ("Capex % revenue", a.capex_pct, PCT),
        ("DSO (days)", a.dso, "0"), ("DIO (days, on COGS)", a.dio, "0"), ("DPO (days, on COGS)", a.dpo, "0"),
        ("Tax rate", a.tax_rate, PCT), ("Debt repayment per year", a.debt_repayment, N0),
        ("Interest rate on opening debt", a.interest_rate, PCT), ("Dividend payout", a.payout, PCT),
        ("WACC", a.wacc, PCT), ("Terminal growth", a.terminal_growth, PCT), ("Shares (millions)", a.shares_m, "0"),
    ]
    for r, (label, v, fmt) in enumerate(scalars, start=8):
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=v)
        c.font, c.number_format = BLUE, fmt
    ws.column_dimensions["A"].width = 34

    # Historicals (row 7 is FY2025)
    wh = wb.create_sheet("Historicals")
    for j, col in enumerate(hist.columns, start=1):
        c = wh.cell(row=1, column=j, value=col)
        c.font, c.fill = BOLD, HEAD
    for i, row in enumerate(hist.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            c = wh.cell(row=i, column=j, value=float(v) if j > 1 else int(v))
            c.font = BLUE
            if j > 1:
                c.number_format = N0
    for j in range(1, len(hist.columns) + 1):
        wh.column_dimensions[get_column_letter(j)].width = 16

    # Model
    wm = wb.create_sheet("Model")
    wm["A1"] = "Projections (formulas)"
    wm["A1"].font = BOLD
    header(wm, 3, YEARS)
    labels = {
        4: "Revenue", 5: "COGS", 6: "SG&A", 7: "EBITDA", 8: "Depreciation", 9: "EBIT", 10: "Interest",
        11: "Profit before tax", 12: "Tax", 13: "Net income", 14: "Receivables", 15: "Inventory",
        16: "Payables", 17: "Change in NWC", 18: "Capex", 19: "Net PPE", 20: "Long-term debt",
        21: "Dividends", 22: "Operating cash flow", 23: "Cash", 24: "Retained earnings",
        25: "Unlevered free cash flow", 26: "Balance check (should be 0)",
    }
    for r, t in labels.items():
        wm.cell(row=r, column=1, value=t)
    H = "Historicals!"
    A = "Assumptions!"
    for i, col in enumerate(COLS):
        prev = COLS[i - 1] if i else None
        acol = "CDEFG"[i]
        first = i == 0
        f = {
            4: f"={H}B7*(1+{A}{acol}4)" if first else f"={prev}4*(1+{A}{acol}4)",
            5: f"={col}4*{A}{acol}5",
            6: f"={col}4*{A}{acol}6",
            7: f"={col}4-{col}5-{col}6",
            8: f"={A}$B$8*{H}L7" if first else f"={A}$B$8*{prev}19",
            9: f"={col}7-{col}8",
            10: f"={A}$B$15*{H}O7" if first else f"={A}$B$15*{prev}20",
            11: f"={col}9-{col}10",
            12: f"={A}$B$13*{col}11",
            13: f"={col}11-{col}12",
            14: f"={col}4*{A}$B$10/365",
            15: f"={col}5*{A}$B$11/365",
            16: f"={col}5*{A}$B$12/365",
            17: (f"=({col}14+{col}15-{col}16)-({H}J7+{H}K7-{H}N7)" if first
                 else f"=({col}14+{col}15-{col}16)-({prev}14+{prev}15-{prev}16)"),
            18: f"={A}$B$9*{col}4",
            19: f"={H}L7+{col}18-{col}8" if first else f"={prev}19+{col}18-{col}8",
            20: f"={H}O7-{A}$B$14" if first else f"={prev}20-{A}$B$14",
            21: f"={A}$B$16*{col}13",
            22: f"={col}13+{col}8-{col}17",
            23: (f"={H}I7+{col}22-{col}18-{A}$B$14-{col}21" if first
                 else f"={prev}23+{col}22-{col}18-{A}$B$14-{col}21"),
            24: f"={H}Q7+{col}13-{col}21" if first else f"={prev}24+{col}13-{col}21",
            25: f"={col}9*(1-{A}$B$13)+{col}8-{col}18-{col}17",
            26: f"=({col}23+{col}14+{col}15+{col}19+{H}M7)-({col}16+{col}20+{H}P7+{col}24)",
        }
        for r, formula in f.items():
            c = wm[f"{col}{r}"]
            c.value = formula
            c.number_format = N0 if r != 26 else "0.00"
    wm.column_dimensions["A"].width = 32

    # DCF
    wd = wb.create_sheet("DCF")
    wd["A1"] = "Discounted cash flow (end of year discounting)"
    wd["A1"].font = BOLD
    wd["A3"], wd["A4"], wd["A5"], wd["A6"] = "Period", "Unlevered FCF", "Discount factor", "PV of FCF"
    for i, col in enumerate(COLS):
        wd[f"{col}3"] = i + 1
        wd[f"{col}4"] = f"=Model!{col}25"
        wd[f"{col}5"] = f"=1/(1+{A}$B$17)^{col}3"
        wd[f"{col}6"] = f"={col}4*{col}5"
        wd[f"{col}4"].number_format = wd[f"{col}6"].number_format = N0
        wd[f"{col}5"].number_format = "0.0000"
    out = [
        (8, "Sum of PV of FCF", "=SUM(B6:F6)", N0),
        (9, "Terminal value", f"=F4*(1+{A}B18)/({A}B17-{A}B18)", N0),
        (10, "PV of terminal value", "=B9*F5", N0),
        (11, "Enterprise value", "=B8+B10", N0),
        (12, "Net debt (debt less cash, FY2025)", f"={H}O7-{H}I7", N0),
        (13, "Equity value", "=B11-B12", N0),
        (14, "Value per share (KES)", f"=B13/{A}B19", "0.00"),
        (15, "Terminal value share of EV", "=B10/B11", PCT),
    ]
    for r, label, formula, fmt in out:
        wd.cell(row=r, column=1, value=label)
        c = wd.cell(row=r, column=2, value=formula)
        c.number_format = fmt
    wd.column_dimensions["A"].width = 34

    # Sensitivity
    wsn = wb.create_sheet("Sensitivity")
    wsn["A1"] = "Enterprise value: WACC (down) against terminal growth (across)"
    wsn["A1"].font = BOLD
    gs, waccs = [0.04, 0.045, 0.05, 0.055, 0.06], [0.11, 0.12, 0.13, 0.14, 0.15]
    for i, g in enumerate(gs):
        c = wsn.cell(row=3, column=2 + i, value=g)
        c.number_format, c.font = PCT, BLUE
    for r, w in enumerate(waccs, start=4):
        c = wsn.cell(row=r, column=1, value=w)
        c.number_format, c.font = PCT, BLUE
        for i in range(5):
            col = COLS[i]
            c = wsn.cell(
                row=r, column=2 + i,
                value=(f"=SUMPRODUCT(DCF!$B$4:$F$4/(1+$A{r})^DCF!$B$3:$F$3)"
                       f"+DCF!$F$4*(1+{col}$3)/($A{r}-{col}$3)/(1+$A{r})^DCF!$F$3"),
            )
            c.number_format = N0

    wb.save(path)


if __name__ == "__main__":
    target = ROOT / "outputs" / "savanna_foods_model.xlsx"
    target.parent.mkdir(exist_ok=True)
    build(target)
    print("wrote", target)
